"""
excel_exporter.py — Export seating plans to Excel (.xlsx) files.

Uses the `openpyxl` library to create Excel workbooks.
Replaces the SheetJS (XLSX.js) export in the original web app.

Three export modes (matching the original):
  1. Single hall  → one sheet with the classroom-style layout
  2. All halls    → one workbook with a sheet per hall
  3. Master list  → one sheet with a sorted table of every student
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from app.constants import YEARS
from app.allotment_engine import build_master_list


# ---------------------------------------------------------------------------
# HELPERS
# ---------------------------------------------------------------------------

def _apply_header_style(ws, row_num: int, max_col: int):
    """Make a header row bold, centered, and merged across all columns."""
    ws.merge_cells(
        start_row=row_num, start_column=1,
        end_row=row_num, end_column=max_col,
    )
    cell = ws.cell(row=row_num, column=1)
    cell.font = Font(name="Calibri", size=12, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _set_column_widths(ws, widths: list[int]):
    """Set the width of each column from a list of integers."""
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


# ---------------------------------------------------------------------------
# HALL LAYOUT SHEET BUILDER
# ---------------------------------------------------------------------------

def _build_hall_sheet(ws, hall: dict, college_name: str, exam_title: str):
    """
    Populate a worksheet with the classroom-style seating layout for one hall.

    Layout matches the original web Excel export:
      Row 1: College name (merged)
      Row 2: Exam title (merged)
      Row 3: "SEATING PLAN: HALL X" (merged)
      Row 4: blank spacer
      Row 5+: Seat data — 6 pairs of columns (roll + seat#) with gaps
    """
    total_cols = 17  # 6 seats × 2 columns + 5 gap columns

    # Header rows
    ws.append([college_name.upper()])
    ws.append([exam_title])
    ws.append([f"SEATING PLAN: HALL {hall['name']}"])
    ws.append([])  # spacer

    _apply_header_style(ws, 1, total_cols)
    _apply_header_style(ws, 2, total_cols)
    _apply_header_style(ws, 3, total_cols)

    # Seat data rows
    for row in hall["rows"]:
        cells = []

        # Left desk (3 seats) — each seat writes roll + seat_no + gap
        for seat in row["left_desk"]:
            cells.append(seat["roll"] if seat else "-")
            cells.append(seat["seat_no"] if seat else "-")
            cells.append("")  # gap column

        # Right desk (3 seats)
        for i, seat in enumerate(row["right_desk"]):
            cells.append(seat["roll"] if seat else "-")
            cells.append(seat["seat_no"] if seat else "-")
            if i < 2:  # gap after first two seats, not the last
                cells.append("")

        ws.append(cells)

    # Column widths: [roll=16, seat=7, gap=4] × 6 seats
    _set_column_widths(ws, [
        16, 7, 4, 16, 7, 4, 16, 7, 7,  # left desk + central gap
        16, 7, 4, 16, 7, 4, 16, 7,      # right desk
    ])


# ---------------------------------------------------------------------------
# PUBLIC API
# ---------------------------------------------------------------------------

def export_single_hall(
    hall: dict,
    college_name: str,
    exam_title: str,
    filepath: str,
) -> None:
    """
    Export a single hall's seating layout to an Excel file.

    Parameters
    ----------
    hall         : one hall dict from the allotment engine
    college_name : e.g. "ANANDA COLLEGE OF ARTS & SCIENCE"
    exam_title   : e.g. "Semester End Examinations - April 2026"
    filepath     : full path for the output .xlsx file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f"Hall {hall['name']}"

    _build_hall_sheet(ws, hall, college_name, exam_title)

    wb.save(filepath)


def export_all_halls(
    halls: list[dict],
    college_name: str,
    exam_title: str,
    filepath: str,
) -> None:
    """
    Export all halls into a single workbook with one sheet per hall.
    """
    wb = Workbook()
    # Remove the default empty sheet created by openpyxl.
    wb.remove(wb.active)

    for hall in halls:
        ws = wb.create_sheet(title=f"Hall {hall['name']}")
        _build_hall_sheet(ws, hall, college_name, exam_title)

    wb.save(filepath)


def export_master_list(
    halls: list[dict],
    college_name: str,
    exam_title: str,
    filepath: str,
) -> None:
    """
    Export a sorted master list of every student allocation to Excel.

    The output is a clean table suitable for printing on notice boards:
      S.No | Register Number | Academic Year | Allocated Hall | Seat Location
    """
    records = build_master_list(halls)

    wb = Workbook()
    ws = wb.active
    ws.title = "Master Allotment List"

    total_cols = 5

    # Header rows
    ws.append([college_name.upper()])
    ws.append([exam_title])
    ws.append(["MASTER SEATING ALLOTMENT LIST (NOTICE BOARD COPY)"])
    ws.append([])  # spacer

    _apply_header_style(ws, 1, total_cols)
    _apply_header_style(ws, 2, total_cols)
    _apply_header_style(ws, 3, total_cols)

    # Table header row
    header_font = Font(name="Calibri", size=10, bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    headers = ["S.No", "Register Number", "Academic Year", "Allocated Hall", "Seat Location"]
    ws.append(headers)
    for col in range(1, total_cols + 1):
        cell = ws.cell(row=5, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for i, record in enumerate(records, start=1):
        year_label = YEARS.get(record["year"], {}).get("label", f"Year {record['year']}")
        ws.append([
            i,
            record["roll"],
            year_label,
            record["hall"],
            record["location"],
        ])

    # Column widths
    _set_column_widths(ws, [8, 22, 18, 18, 38])

    wb.save(filepath)
